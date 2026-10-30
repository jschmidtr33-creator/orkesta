"""Lectura segura de CSV y XLSX.

- CSV: detecta la codificación (UTF-8 con o sin BOM, Windows-1252, Latin-1) y el delimitador
  (coma, punto y coma, tabulador o barra vertical).
- XLSX: lee la primera hoja en modo solo lectura y rechaza archivos comprimidos anómalos.
- Límites de tamaño, filas, columnas y longitud de celda para que un archivo malformado o
  malicioso no agote la memoria del servicio.
"""

from __future__ import annotations

import csv
import io
import zipfile
from dataclasses import dataclass
from datetime import date, datetime
from typing import Literal

from orkesta.engine.normalize.text import clean

FileFormat = Literal["csv", "xlsx"]

MAX_COLUMNS = 200
MAX_CELL_CHARS = 2000
# Un XLSX es un ZIP: se limita el tamaño descomprimido para evitar "bombas zip".
MAX_XLSX_UNCOMPRESSED_BYTES = 200 * 1024 * 1024
_ENCODINGS = ("utf-8-sig", "cp1252", "latin-1")
_DELIMITERS = ",;\t|"


class ImportFileError(ValueError):
    """El archivo no se puede importar. El mensaje está pensado para mostrarse al usuario."""


@dataclass(frozen=True)
class RawTable:
    headers: list[str]
    rows: list[list[str]]
    file_format: FileFormat
    encoding: str | None = None
    delimiter: str | None = None


@dataclass(frozen=True)
class ReadLimits:
    max_bytes: int
    max_rows: int


def detect_format(filename: str) -> FileFormat:
    name = filename.lower().strip()
    if name.endswith((".csv", ".txt", ".tsv")):
        return "csv"
    if name.endswith(".xlsx"):
        return "xlsx"
    if name.endswith(".xls"):
        raise ImportFileError(
            "El formato .xls (Excel 97-2003) no es compatible. Guárdalo como .xlsx o .csv."
        )
    raise ImportFileError("Formato no compatible. Sube un archivo .csv o .xlsx.")


def _decode(content: bytes) -> tuple[str, str]:
    for encoding in _ENCODINGS:
        try:
            return content.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    raise ImportFileError("No se pudo leer la codificación del archivo.")  # pragma: no cover


def _detect_delimiter(sample: str) -> str:
    try:
        return csv.Sniffer().sniff(sample, delimiters=_DELIMITERS).delimiter
    except csv.Error:
        first_line = sample.splitlines()[0] if sample else ""
        counts = {d: first_line.count(d) for d in _DELIMITERS}
        best = max(counts, key=lambda d: counts[d])
        return best if counts[best] > 0 else ","


def _cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        text = value.isoformat(sep=" ")
    elif isinstance(value, date):
        text = value.isoformat()
    elif isinstance(value, float) and value.is_integer():
        text = str(int(value))
    else:
        text = str(value)
    text = clean(text)
    if len(text) > MAX_CELL_CHARS:
        raise ImportFileError(f"Una celda supera el máximo de {MAX_CELL_CHARS} caracteres.")
    return text


def _finalize(raw_rows: list[list[str]], limits: ReadLimits, **meta: object) -> RawTable:
    rows = [row for row in raw_rows if any(cell for cell in row)]
    if not rows:
        raise ImportFileError("El archivo está vacío.")
    header_row, data = rows[0], rows[1:]
    if len(header_row) > MAX_COLUMNS:
        raise ImportFileError(f"El archivo tiene más de {MAX_COLUMNS} columnas.")
    if not data:
        raise ImportFileError("El archivo solo tiene encabezados, sin filas de datos.")
    if len(data) > limits.max_rows:
        raise ImportFileError(
            f"El archivo tiene {len(data):,} filas; el máximo por lote es {limits.max_rows:,}."
        )

    headers: list[str] = []
    seen: dict[str, int] = {}
    for index, name in enumerate(header_row, start=1):
        label = name or f"Columna {index}"
        if label in seen:
            seen[label] += 1
            label = f"{label} ({seen[label]})"
        else:
            seen[label] = 1
        headers.append(label)

    width = len(headers)
    normalized = []
    for row in data:
        if len(row) > width and any(row[width:]):
            raise ImportFileError(
                "Algunas filas tienen más valores que encabezados. Revisa el delimitador o las "
                "comillas del archivo."
            )
        normalized.append((row + [""] * width)[:width])
    return RawTable(headers=headers, rows=normalized, **meta)  # type: ignore[arg-type]


def read_csv(content: bytes, limits: ReadLimits) -> RawTable:
    text, encoding = _decode(content)
    delimiter = _detect_delimiter(text[:20000])
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)
    rows: list[list[str]] = []
    try:
        for record in reader:
            if len(rows) > limits.max_rows + 1:
                break
            if len(record) > MAX_COLUMNS:
                raise ImportFileError(f"El archivo tiene más de {MAX_COLUMNS} columnas.")
            rows.append([_cell(value) for value in record])
    except csv.Error as exc:
        raise ImportFileError(f"El CSV está mal formado: {exc}") from exc
    return _finalize(rows, limits, file_format="csv", encoding=encoding, delimiter=delimiter)


def read_xlsx(content: bytes, limits: ReadLimits) -> RawTable:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            uncompressed = sum(info.file_size for info in archive.infolist())
    except zipfile.BadZipFile as exc:
        raise ImportFileError("El archivo .xlsx está dañado o no es un Excel válido.") from exc
    if uncompressed > MAX_XLSX_UNCOMPRESSED_BYTES:
        raise ImportFileError("El archivo .xlsx es demasiado grande una vez descomprimido.")

    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl lanza tipos variados ante archivos dañados
        raise ImportFileError("El archivo .xlsx está dañado o no es un Excel válido.") from exc
    try:
        sheet = workbook.worksheets[0]
        rows: list[list[str]] = []
        for record in sheet.iter_rows(values_only=True):
            if len(rows) > limits.max_rows + 1:
                break
            values = list(record)
            while values and values[-1] in (None, ""):
                values.pop()
            if len(values) > MAX_COLUMNS:
                raise ImportFileError(f"El archivo tiene más de {MAX_COLUMNS} columnas.")
            rows.append([_cell(v) for v in values])
    finally:
        workbook.close()
    width = max((len(r) for r in rows), default=0)
    rows = [r + [""] * (width - len(r)) for r in rows]
    return _finalize(rows, limits, file_format="xlsx")


def read_table(content: bytes, filename: str, limits: ReadLimits) -> RawTable:
    if len(content) == 0:
        raise ImportFileError("El archivo está vacío.")
    if len(content) > limits.max_bytes:
        megabytes = limits.max_bytes / (1024 * 1024)
        raise ImportFileError(f"El archivo supera el máximo de {megabytes:.0f} MB.")
    file_format = detect_format(filename)
    if file_format == "xlsx":
        return read_xlsx(content, limits)
    return read_csv(content, limits)
