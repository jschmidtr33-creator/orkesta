"""Lectura de bases de contactos (Excel con todas sus hojas, o CSV) y paso a registros limpios.

Solo se conservan 8 campos: nombre y apellido, ubicación (ciudad y país), teléfono, correo,
cargo, empresa e industria. El resto de columnas del archivo se ignora. Las columnas se
reconocen por su encabezado o, si no lo hay, por su contenido.
"""

from __future__ import annotations

import csv
import io
import re
import zipfile
from collections.abc import Callable, Iterator
from dataclasses import dataclass

from orkesta.engine.files.reader import (
    MAX_CELL_CHARS,
    MAX_COLUMNS,
    MAX_XLSX_UNCOMPRESSED_BYTES,
    ImportFileError,
    _decode,
    _detect_delimiter,
    detect_format,
)
from orkesta.engine.files.reader import (
    _cell as _reader_cell,
)
from orkesta.engine.net.domains import registrable_domain
from orkesta.engine.normalize import normalize_company, parse_location
from orkesta.engine.normalize.names import split_full_name
from orkesta.engine.normalize.text import clean, fold
from orkesta.engine.taxonomy.industries import classify

MAX_BYTES = 30 * 1024 * 1024
MAX_ROWS = 200_000

# Correo gratuito: es un dato de contacto válido, pero no enseña el patrón de una empresa.
FREEMAIL_DOMAINS = frozenset(
    [
        "gmail.com",
        "hotmail.com",
        "outlook.com",
        "yahoo.com",
        "yahoo.com.mx",
        "yahoo.es",
        "hotmail.es",
        "live.com",
        "live.com.mx",
        "msn.com",
        "icloud.com",
        "me.com",
        "aol.com",
        "prodigy.net.mx",
        "outlook.es",
        "hotmail.com.ar",
        "yahoo.com.ar",
        "gmail.co",
        "protonmail.com",
        "proton.me",
        "terra.com.br",
        "uol.com.br",
        "bol.com.br",
        "hotmail.cl",
        "yahoo.cl",
        "live.cl",
        "gmx.com",
        "mail.com",
        "zoho.com",
        "yandex.com",
        "outlook.com.mx",
        "hotmail.com.mx",
    ]
)

# Campo → nombres de columna reconocidos (plegados: sin tildes ni mayúsculas).
FIELD_ALIASES: dict[str, tuple[str, ...]] = {
    "first": ("first_name", "first name", "nombre", "nombres"),
    "last": ("last_name", "last name", "apellido", "apellidos"),
    "full": ("full_name", "full name", "nombre completo", "nombre y apellido", "nombre"),
    "email": (
        "correo electronico", "email", "e-mail", "correo", "mail", "work email",
        "direccion de correo electronico",
    ),
    "title": ("title", "cargo", "puesto", "job title", "position"),
    "company": (
        "company", "empresa", "empresa organizacion", "empresa  organizacion", "organizacion",
        "company name", "compania",
    ),
    "phone": (
        "mobile_phone1", "phone_number1", "telefono", "numero telefono movil", "celular",
        "movil", "phone", "telefono movil",
    ),
    "city": ("locality", "ciudad", "city"),
    "country": ("country", "pais", "country_code"),
    "location": ("location", "ubicacion"),
    "industry": ("company_industry", "industria", "industry", "sector"),
}  # fmt: skip

# Prefijos telefónicos y extensiones de país para completar el país cuando falta.
PHONE_PREFIXES: tuple[tuple[str, str], ...] = (
    ("52", "MX"), ("56", "CL"), ("57", "CO"), ("51", "PE"), ("54", "AR"), ("593", "EC"),
    ("598", "UY"), ("506", "CR"), ("502", "GT"), ("507", "PA"), ("591", "BO"), ("595", "PY"),
    ("503", "SV"), ("504", "HN"), ("505", "NI"), ("58", "VE"), ("55", "BR"), ("34", "ES"),
)  # fmt: skip
TLD_COUNTRIES = {
    "mx": "MX", "cl": "CL", "co": "CO", "pe": "PE", "ar": "AR", "ec": "EC", "uy": "UY",
    "cr": "CR", "gt": "GT", "pa": "PA", "do": "DO", "bo": "BO", "py": "PY", "es": "ES",
    "br": "BR", "ve": "VE", "hn": "HN", "sv": "SV", "ni": "NI",
}  # fmt: skip


def _cell(value: object) -> str:
    """Celda como texto (como el lector de lotes), pero las muy largas se recortan en vez de
    rechazar el archivo: suelen ser descripciones de empresa, que no se guardan."""
    if isinstance(value, str) and len(value) > MAX_CELL_CHARS:
        value = value[:MAX_CELL_CHARS]
    return _reader_cell(value)


_EMAIL = re.compile(r"^[a-z0-9][a-z0-9._%+'-]{0,63}@[a-z0-9-]+(\.[a-z0-9-]+)+$")
_PHONE_CHARS = re.compile(r"[^\d+]")


@dataclass(frozen=True)
class Sheet:
    name: str
    headers: list[str]
    rows: list[list[str]]


@dataclass(frozen=True)
class ProspectRecord:
    full_name: str
    first_name: str
    last_name: str
    title: str | None
    company: str
    company_key: str
    email: str | None
    email_domain: str | None
    phone: str | None
    city: str | None
    country_code: str | None
    industry: str | None
    industry_code: str | None

    @property
    def dedup_key(self) -> str:
        if self.email:
            return self.email
        return f"{fold(self.full_name)}|{self.company_key}"[:400]

    @property
    def corporate_email(self) -> bool:
        return bool(self.email_domain) and self.email_domain not in FREEMAIL_DOMAINS


# --- Lectura ---------------------------------------------------------------------------------


def read_sheets(content: bytes, filename: str) -> list[Sheet]:
    """Todas las hojas con datos (Excel) o la tabla (CSV)."""
    if not content:
        raise ImportFileError("El archivo está vacío.")
    if len(content) > MAX_BYTES:
        raise ImportFileError(f"El archivo supera el máximo de {MAX_BYTES // (1024 * 1024)} MB.")
    if detect_format(filename) == "xlsx":
        raw = list(_xlsx_tables(content))
    else:
        text, _ = _decode(content)
        delimiter = _detect_delimiter(text[:20000])
        try:
            rows = [
                [_cell(v) for v in r] for r in csv.reader(io.StringIO(text), delimiter=delimiter)
            ]
        except csv.Error as exc:
            raise ImportFileError(f"El CSV está mal formado: {exc}") from exc
        raw = [("csv", rows)]
    sheets: list[Sheet] = []
    total = 0
    for name, rows in raw:
        rows = [r for r in rows if any(c for c in r)]
        if len(rows) < 2 and not (rows and _looks_headerless(rows[0])):
            continue
        width = max(len(r) for r in rows)
        if width > MAX_COLUMNS:
            raise ImportFileError(f"La hoja {name} tiene más de {MAX_COLUMNS} columnas.")
        rows = [r + [""] * (width - len(r)) for r in rows]
        if _looks_headerless(rows[0]):
            headers, body = [""] * width, rows
        else:
            headers, body = rows[0], rows[1:]
        total += len(body)
        if total > MAX_ROWS:
            raise ImportFileError(f"El archivo supera {MAX_ROWS:,} filas.")
        sheets.append(Sheet(name, headers, body))
    if not sheets:
        raise ImportFileError("El archivo no tiene filas de datos.")
    return sheets


def _xlsx_tables(content: bytes) -> Iterator[tuple[str, list[list[str]]]]:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(i.file_size for i in archive.infolist()) > MAX_XLSX_UNCOMPRESSED_BYTES:
                raise ImportFileError("El archivo .xlsx es demasiado grande descomprimido.")
    except zipfile.BadZipFile as exc:
        raise ImportFileError("El archivo .xlsx está dañado o no es un Excel válido.") from exc
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl lanza tipos variados ante archivos dañados
        raise ImportFileError("El archivo .xlsx está dañado o no es un Excel válido.") from exc
    try:
        for sheet in workbook.worksheets:
            rows = [[_cell(v) for v in r] for r in sheet.iter_rows(values_only=True)]
            yield sheet.title, rows
    finally:
        workbook.close()


def _looks_headerless(row: list[str]) -> bool:
    """Una primera fila con un correo es un dato, no un encabezado."""
    return any("@" in c and _EMAIL.match(c.strip().lower()) for c in row)


# --- Columnas --------------------------------------------------------------------------------


def map_columns(sheet: Sheet) -> dict[str, int]:
    """Campo → índice de columna. Por encabezado; si no hay encabezados, por contenido."""
    if any(sheet.headers):
        return _map_by_headers(sheet.headers)
    return _map_by_content(sheet.rows)


def _canonical(text: str) -> str:
    """ "E-mail", "e_mail" y "E mail" son el mismo encabezado."""
    return re.sub(r"[^a-z0-9]", "", fold(text))


def _map_by_headers(headers: list[str]) -> dict[str, int]:
    folded = [_canonical(h) for h in headers]
    out: dict[str, int] = {}
    for field, names in FIELD_ALIASES.items():
        for name in names:
            target = _canonical(name)
            index = next((i for i, h in enumerate(folded) if h == target), None)
            if index is not None and index not in out.values():
                out[field] = index
                break
    # "Nombre" es nombre de pila cuando hay apellido; si no, es el nombre completo.
    if "last" in out:
        out.pop("full", None)
    elif "first" in out and "full" not in out:
        out["full"] = out.pop("first")
    else:
        out.pop("first", None)
    return out


def _map_by_content(rows: list[list[str]]) -> dict[str, int]:
    """Archivos sin encabezados: correo, teléfono, país, industria y cargo por su contenido;
    las dos primeras columnas de texto son nombre y apellido, y la siguiente libre, la empresa."""
    sample = rows[:300]
    width = len(sample[0]) if sample else 0

    def share(index: int, test: Callable[[str], bool]) -> float:
        values = [r[index] for r in sample if r[index]]
        return sum(1 for v in values if test(v)) / max(len(values), 1)

    out: dict[str, int] = {}
    scores: dict[str, Callable[[str], bool]] = {
        "email": lambda v: bool(_EMAIL.match(v.strip().lower())),
        "phone": lambda v: (
            7 <= len(re.sub(r"\D", "", v)) <= 15 and not re.search(r"[a-z]", v, re.I)
        ),
        "country": lambda v: parse_location(v).country_code is not None and len(v) < 40,
        "industry": lambda v: classify(v) is not None,
        "title": lambda v: bool(set(fold(v).split()) & _TITLE_WORDS),
    }

    def repetition(index: int) -> float:
        values = [r[index] for r in sample if r[index]]
        return 1 - len(set(values)) / max(len(values), 1)

    for field, test in scores.items():
        # A igual proporción, la columna más repetitiva: la industria es una categoría que se
        # repite; los nombres de empresa varían aunque alguno parezca industria ("Retail X").
        best = max(
            (i for i in range(width) if i not in out.values()),
            key=lambda i: (round(share(i, test), 2), repetition(i), i),
            default=None,
        )
        if best is not None and share(best, test) >= 0.6:
            out[field] = best
    free = [i for i in range(width) if i not in out.values()]
    for field, index in zip(("first", "last", "company"), free, strict=False):
        out[field] = index
    return out


_TITLE_WORDS = frozenset(
    [
        "gerente",
        "director",
        "directora",
        "jefe",
        "jefa",
        "coordinador",
        "coordinadora",
        "analista",
        "ejecutivo",
        "ejecutiva",
        "manager",
        "head",
        "chief",
        "officer",
        "vp",
        "vicepresidente",
        "presidente",
        "ceo",
        "cfo",
        "cto",
        "cio",
        "coo",
        "cmo",
        "lider",
        "lead",
        "especialista",
        "specialist",
        "subdirector",
        "subgerente",
        "supervisor",
        "ingeniero",
        "asesor",
        "consultor",
        "socio",
        "partner",
        "fundador",
        "founder",
        "owner",
        "dueno",
        "administrador",
    ]
)


# --- Registros -------------------------------------------------------------------------------


def _value(row: list[str], columns: dict[str, int], field: str) -> str:
    index = columns.get(field)
    return clean(row[index]) if index is not None and index < len(row) else ""


def _phone(value: str) -> str | None:
    digits = _PHONE_CHARS.sub("", value)
    return digits[:40] if len(re.sub(r"\D", "", digits)) >= 7 else None


def infer_country(phone: str | None, email_domain: str | None) -> str | None:
    if phone and phone.startswith("+"):
        number = phone[1:]
        for prefix, country in sorted(PHONE_PREFIXES, key=lambda p: -len(p[0])):
            if number.startswith(prefix):
                return country
    if email_domain:
        return TLD_COUNTRIES.get(email_domain.rsplit(".", 1)[-1])
    return None


def to_records(sheet: Sheet, columns: dict[str, int]) -> list[ProspectRecord]:
    records: list[ProspectRecord] = []
    for row in sheet.rows:
        first, last = _value(row, columns, "first"), _value(row, columns, "last")
        if not (first and last):
            full = _value(row, columns, "full")
            if full:
                split = split_full_name(full)
                first, last = split.first or first, split.last or last
        company = _value(row, columns, "company")
        if not first or not company:
            continue
        email = _value(row, columns, "email").lower().strip(" <>;,") or None
        if email and not _EMAIL.match(email):
            email = None
        email_domain = None
        if email:
            host = email.split("@", 1)[1]
            email_domain = registrable_domain(host) or host
        phone = _phone(_value(row, columns, "phone"))
        location = parse_location(
            ", ".join(
                v for v in (_value(row, columns, "city"), _value(row, columns, "country")) if v
            )
            or _value(row, columns, "location")
        )
        industry = _value(row, columns, "industry") or None
        company_name = normalize_company(company)
        records.append(
            ProspectRecord(
                full_name=" ".join(x for x in (first, last) if x)[:300],
                first_name=first[:200],
                last_name=last[:200],
                title=_value(row, columns, "title")[:300] or None,
                company=company[:300],
                company_key=(company_name.match_key or fold(company))[:300],
                email=email,
                email_domain=email_domain,
                phone=phone,
                city=(_value(row, columns, "city") or location.city or "")[:120] or None,
                country_code=location.country_code or infer_country(phone, email_domain),
                industry=industry[:200] if industry else None,
                industry_code=classify(industry) if industry else None,
            )
        )
    return records


def records_from_file(content: bytes, filename: str) -> list[ProspectRecord]:
    """Registros de todas las hojas, sin duplicados (el primero completa con los siguientes)."""
    merged: dict[str, ProspectRecord] = {}
    for sheet in read_sheets(content, filename):
        columns = map_columns(sheet)
        if "company" not in columns or not ({"first", "full"} & set(columns)):
            continue
        for record in to_records(sheet, columns):
            previous = merged.get(record.dedup_key)
            merged[record.dedup_key] = merge(previous, record) if previous else record
    return list(merged.values())


def merge(base: ProspectRecord, extra: ProspectRecord) -> ProspectRecord:
    """Completa los campos vacíos de ``base`` con los de ``extra``."""
    values = {
        field: getattr(base, field) or getattr(extra, field)
        for field in ProspectRecord.__dataclass_fields__
    }
    return ProspectRecord(**values)


# --- Patrones reales por dominio -------------------------------------------------------------
