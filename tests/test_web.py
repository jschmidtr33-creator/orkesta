"""Recorrido de la interfaz web con PostgreSQL real: acceso, importación, exportación, borrado."""

from __future__ import annotations

import csv
import io
import re
import uuid
from collections.abc import Iterator
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import Engine
from sqlalchemy.orm import Session, sessionmaker

from orkesta.api.main import create_app
from orkesta.auth import create_user
from orkesta.db.models import Organization
from orkesta.worker.__main__ import FAILURE_HOOKS, HANDLERS
from orkesta.worker.queue import RetryPolicy
from orkesta.worker.runner import Worker

pytestmark = pytest.mark.db

FIXTURE = (Path(__file__).parent / "fixtures" / "contactos_ejemplo.csv").read_bytes()
PASSWORD = "contraseña-de-prueba"


def _make_user(factory: sessionmaker[Session], email: str) -> None:
    with factory.begin() as session:
        org = Organization(name=email, slug=f"org-{uuid.uuid4().hex[:8]}")
        session.add(org)
        session.flush()
        create_user(session, email=email, password=PASSWORD, organization_id=org.id)


def _csrf(client: TestClient, path: str = "/login") -> str:
    match = re.search(r'name="csrf_token" value="([^"]+)"', client.get(path).text)
    assert match, f"no hay token CSRF en {path}"
    return match.group(1)


def _login(client: TestClient, email: str, password: str = PASSWORD) -> None:
    token = _csrf(client)
    response = client.post(
        "/login", data={"email": email, "password": password, "csrf_token": token}
    )
    assert response.status_code == 303, response.text


@pytest.fixture()
def client(engine: Engine, session_factory: sessionmaker[Session]) -> Iterator[TestClient]:
    _make_user(session_factory, "jenn@ejemplo.com")
    with TestClient(create_app(engine_provider=lambda: engine), follow_redirects=False) as c:
        yield c


def _upload(client: TestClient, content: bytes = FIXTURE, name: str = "contactos.csv") -> str:
    response = client.post(
        "/lotes",
        data={"csrf_token": _csrf(client, "/lotes/nuevo")},
        files={"archivo": (name, content, "text/csv")},
    )
    assert response.status_code == 303, response.text
    match = re.match(r"/lotes/([0-9a-f-]+)/mapeo", response.headers["location"])
    assert match
    return match.group(1)


def _confirm(client: TestClient, batch_id: str) -> None:
    response = client.post(
        f"/lotes/{batch_id}/confirmar",
        data={
            "csrf_token": _csrf(client, f"/lotes/{batch_id}/mapeo"),
            "first_name": "Nombre",
            "last_name": "Apellido",
            "title": "Cargo",
            "company": "Empresa",
            "location": "Ubicación",
        },
    )
    assert response.status_code == 303
    assert response.headers["location"] == f"/lotes/{batch_id}"


def _run_worker(factory: sessionmaker[Session]) -> None:
    worker = Worker(
        factory,
        HANDLERS,
        retry_policy=RetryPolicy(1, 1),
        poll_interval_seconds=0.01,
        lock_timeout_seconds=600,
        failure_hooks=FAILURE_HOOKS,
    )
    assert worker.run_once() is True


def test_pages_require_login(client: TestClient) -> None:
    response = client.get("/lotes")
    assert response.status_code == 303
    assert response.headers["location"] == "/login?next=/lotes"


def test_login_rejects_missing_csrf(client: TestClient) -> None:
    response = client.post("/login", data={"email": "jenn@ejemplo.com", "password": PASSWORD})
    assert response.status_code == 403


def test_wrong_password_then_lockout(client: TestClient) -> None:
    for _ in range(5):
        token = _csrf(client)
        response = client.post(
            "/login", data={"email": "jenn@ejemplo.com", "password": "mala", "csrf_token": token}
        )
        assert response.status_code == 401
        assert "Correo o contraseña incorrectos" in response.text
    token = _csrf(client)
    response = client.post(
        "/login", data={"email": "jenn@ejemplo.com", "password": PASSWORD, "csrf_token": token}
    )
    assert response.status_code == 429


def test_login_ignores_external_next(client: TestClient) -> None:
    token = _csrf(client)
    response = client.post(
        "/login",
        data={
            "email": "jenn@ejemplo.com",
            "password": PASSWORD,
            "csrf_token": token,
            "next": "//sitio-malicioso.invalid",
        },
    )
    assert response.headers["location"] == "/lotes"


def test_security_headers(client: TestClient) -> None:
    response = client.get("/login")
    assert response.headers["x-frame-options"] == "DENY"
    assert "default-src 'self'" in response.headers["content-security-policy"]
    assert response.headers["x-content-type-options"] == "nosniff"


def test_full_import_flow(client: TestClient, session_factory: sessionmaker[Session]) -> None:
    _login(client, "jenn@ejemplo.com")
    assert "Todavía no hay lotes" in client.get("/lotes").text

    batch_id = _upload(client)
    mapping = client.get(f"/lotes/{batch_id}/mapeo")
    assert mapping.status_code == 200
    assert "Resultado previsto" in mapping.text
    assert "Bancolombia" in mapping.text
    assert 'value="Empresa" selected' in mapping.text

    _confirm(client, batch_id)
    importing = client.get(f"/lotes/{batch_id}")
    assert "Importando" in importing.text
    assert 'http-equiv="refresh"' in importing.text

    _run_worker(session_factory)
    detail = client.get(f"/lotes/{batch_id}")
    assert "Importado" in detail.text and 'http-equiv="refresh"' not in detail.text
    assert "Contactos nuevos" in detail.text

    errors_tab = client.get(f"/lotes/{batch_id}?estado=invalid")
    assert "La empresa no identifica a una organización" in errors_tab.text

    listing = client.get("/lotes")
    assert "contactos.csv" in listing.text

    xlsx = client.get(f"/lotes/{batch_id}/exportar?formato=xlsx&filas=all")
    assert xlsx.status_code == 200
    assert "attachment" in xlsx.headers["content-disposition"]
    sheet = load_workbook(io.BytesIO(xlsx.content)).active
    assert sheet is not None and sheet.max_row == 10

    valid_csv = client.get(f"/lotes/{batch_id}/exportar?formato=csv&filas=valid")
    records = list(csv.reader(io.StringIO(valid_csv.content.decode("utf-8-sig"))))
    assert len(records) == 7 and {r[5] for r in records[1:]} == {"válida"}


def test_mapping_preview_can_change_columns(client: TestClient) -> None:
    _login(client, "jenn@ejemplo.com")
    batch_id = _upload(client)
    response = client.get(
        f"/lotes/{batch_id}/mapeo?first_name=Nombre&last_name=Apellido&company=&location="
    )
    assert "Falta asignar la columna de Empresa." in response.text
    assert "Confirmar importación" not in response.text


def test_invalid_upload_shows_error(client: TestClient) -> None:
    _login(client, "jenn@ejemplo.com")
    response = client.post(
        "/lotes",
        data={"csrf_token": _csrf(client, "/lotes/nuevo")},
        files={"archivo": ("base.xls", b"\xd0\xcf", "application/vnd.ms-excel")},
    )
    assert response.status_code == 400
    assert "Guárdalo como .xlsx o .csv" in response.text


def test_duplicate_upload_warns(client: TestClient) -> None:
    _login(client, "jenn@ejemplo.com")
    _upload(client)
    second = _upload(client)
    assert "Este mismo archivo ya se subió" in client.get(f"/lotes/{second}/mapeo").text


def test_other_organization_cannot_see_batch(
    client: TestClient, session_factory: sessionmaker[Session]
) -> None:
    _login(client, "jenn@ejemplo.com")
    batch_id = _upload(client)
    _make_user(session_factory, "otra@ejemplo.com")
    client.cookies.clear()
    _login(client, "otra@ejemplo.com")
    assert client.get(f"/lotes/{batch_id}/mapeo").status_code == 404
    assert client.get(f"/lotes/{batch_id}/exportar").status_code == 404
    assert "contactos.csv" not in client.get("/lotes").text


def test_delete_batch(client: TestClient, session_factory: sessionmaker[Session]) -> None:
    _login(client, "jenn@ejemplo.com")
    batch_id = _upload(client)
    _confirm(client, batch_id)
    _run_worker(session_factory)
    token = _csrf(client, f"/lotes/{batch_id}/eliminar")
    response = client.post(f"/lotes/{batch_id}/eliminar", data={"csrf_token": token})
    assert response.status_code == 303 and response.headers["location"] == "/lotes"
    assert "eliminado (6 contactos borrados)" in client.get("/lotes").text
    assert client.get(f"/lotes/{batch_id}").status_code == 404


def test_logout(client: TestClient) -> None:
    _login(client, "jenn@ejemplo.com")
    token = _csrf(client, "/lotes")
    assert client.post("/logout", data={"csrf_token": token}).status_code == 303
    assert client.get("/lotes").status_code == 303
