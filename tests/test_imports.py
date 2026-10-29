"""Pruebas de lectura, mapeo, validación y exportación (sin base de datos).

Los contactos de ``tests/fixtures`` son personas ficticias creadas para estas pruebas.
"""

from __future__ import annotations

import csv
import io
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from orkesta.imports.export import ExportRow, sanitize_cell, to_csv, to_xlsx
from orkesta.imports.mapping import propose_mapping, validate_mapping
from orkesta.imports.reader import ImportFileError, ReadLimits, read_table
from orkesta.imports.rows import ISSUES, process_rows

FIXTURE = Path(__file__).parent / "fixtures" / "contactos_ejemplo.csv"
LIMITS = ReadLimits(max_bytes=5 * 1024 * 1024, max_rows=1000)
MESSAGES = {code: message for code, (_, message) in ISSUES.items()}


def _xlsx(rows: list[list[object]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# --- lectura -------------------------------------------------------------------------------


def test_reads_utf8_csv_fixture() -> None:
    table = read_table(FIXTURE.read_bytes(), "contactos.csv", LIMITS)
    assert table.headers == ["Nombre", "Apellido", "Cargo", "Empresa", "Ubicación"]
    assert len(table.rows) == 9
    assert table.rows[0][4] == "Bogotá, Distrito Capital, Colombia"
    assert table.encoding == "utf-8-sig" and table.delimiter == ","


def test_reads_windows_1252_semicolon_csv() -> None:
    content = "Nombre;Apellido;Empresa\nJosé;Muñoz;Compañía Ñandú\n".encode("cp1252")
    table = read_table(content, "x.csv", LIMITS)
    assert table.delimiter == ";" and table.encoding == "cp1252"
    assert table.rows == [["José", "Muñoz", "Compañía Ñandú"]]


def test_reads_tab_separated_with_bom() -> None:
    content = "\ufeffNombre\tApellido\nAna\tLópez\n".encode()
    table = read_table(content, "x.tsv", LIMITS)
    assert table.headers == ["Nombre", "Apellido"]
    assert table.rows == [["Ana", "López"]]


def test_skips_empty_rows_and_pads_short_rows() -> None:
    table = read_table(b"a,b,c\n\n1,2\n,,\n4,5,6\n", "x.csv", LIMITS)
    assert table.rows == [["1", "2", ""], ["4", "5", "6"]]


def test_duplicate_and_empty_headers_get_unique_names() -> None:
    table = read_table(b"Nombre,Nombre,\nA,B,C\n", "x.csv", LIMITS)
    assert table.headers == ["Nombre", "Nombre (2)", "Columna 3"]


def test_reads_xlsx_first_sheet() -> None:
    content = _xlsx([["Nombre", "Apellido", "Teléfono"], ["Ana", "López", 3001234567.0]])
    table = read_table(content, "base.xlsx", LIMITS)
    assert table.file_format == "xlsx"
    assert table.rows == [["Ana", "López", "3001234567"]]


@pytest.mark.parametrize(
    ("content", "filename", "message"),
    [
        (b"", "x.csv", "vacío"),
        (b"a,b\n", "x.csv", "solo tiene encabezados"),
        (b"a,b\n1,2\n", "x.xls", ".xls"),
        (b"a,b\n1,2\n", "x.pdf", "Formato no compatible"),
        (b"not a zip", "x.xlsx", "dañado"),
        (b"a,b\n1,2,3\n", "x.csv", "más valores que encabezados"),
    ],
)
def test_rejects_bad_files(content: bytes, filename: str, message: str) -> None:
    with pytest.raises(ImportFileError, match=message):
        read_table(content, filename, LIMITS)


def test_rejects_too_many_rows() -> None:
    content = ("a\n" + "x\n" * 11).encode()
    with pytest.raises(ImportFileError, match="máximo por lote es 10"):
        read_table(content, "x.csv", ReadLimits(max_bytes=10_000, max_rows=10))


def test_rejects_oversized_file() -> None:
    with pytest.raises(ImportFileError, match="supera el máximo"):
        read_table(b"a\n" + b"x" * 2000, "x.csv", ReadLimits(max_bytes=1000, max_rows=10))


def test_rejects_xlsx_zip_bomb(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr("orkesta.imports.reader.MAX_XLSX_UNCOMPRESSED_BYTES", 1000)
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xl/big.xml", "0" * 100_000)
    with pytest.raises(ImportFileError, match="descomprimido"):
        read_table(buffer.getvalue(), "x.xlsx", LIMITS)


# --- mapeo ---------------------------------------------------------------------------------


def test_proposes_spanish_mapping() -> None:
    proposed = propose_mapping(["Nombre", "Apellido", "Cargo", "Empresa", "Ubicación"])
    assert proposed.mapping == {
        "first_name": "Nombre",
        "last_name": "Apellido",
        "full_name": None,
        "title": "Cargo",
        "company": "Empresa",
        "location": "Ubicación",
    }
    assert {v for k, v in proposed.confidence.items() if k != "full_name"} == {"exact"}


def test_proposes_english_crm_mapping() -> None:
    headers = ["First Name", "Last Name", "Job Title", "Company Name", "Person Location", "Email"]
    proposed = propose_mapping(headers)
    assert proposed.mapping == {
        "first_name": "First Name",
        "last_name": "Last Name",
        "full_name": None,
        "title": "Job Title",
        "company": "Company Name",
        "location": "Person Location",
    }


def test_single_name_column_is_mapped_as_full_name() -> None:
    proposed = propose_mapping(["NOMBRE", "EMAIL", "CARGO", "EMPRESA", "UBICACIÓN"])
    assert proposed.mapping["full_name"] == "NOMBRE"
    assert proposed.mapping["first_name"] is None and proposed.mapping["last_name"] is None
    assert (
        validate_mapping(proposed.mapping, ["NOMBRE", "EMAIL", "CARGO", "EMPRESA", "UBICACIÓN"])
        == []
    )


def test_explicit_full_name_header() -> None:
    proposed = propose_mapping(["Nombre completo", "Empresa"])
    assert proposed.mapping["full_name"] == "Nombre completo"


def test_separate_columns_win_over_full_name() -> None:
    proposed = propose_mapping(["Nombre", "Apellido", "Nombre completo", "Empresa"])
    assert (proposed.mapping["first_name"], proposed.mapping["last_name"]) == ("Nombre", "Apellido")


def test_full_name_rows_are_split_and_flagged() -> None:
    content = (FIXTURE.parent / "contactos_nombre_completo.csv").read_bytes()
    table = read_table(content, "leads.csv", LIMITS)
    rows = [(i, dict(zip(table.headers, r, strict=True))) for i, r in enumerate(table.rows, 2)]
    processed, summary = process_rows(rows, propose_mapping(table.headers).mapping)
    by_row = {r.row_number: r for r in processed}
    assert (summary.valid, summary.invalid) == (7, 1)
    assert (by_row[2].person.first_display, by_row[2].person.last_display) == (
        "Rodrigo Alonso",
        "Quintana Mora",
    )
    assert by_row[3].person.last_display == "Salinas" and "NAME_CLEANED" in by_row[3].issues
    assert "LAST_NAME_INITIAL" in by_row[4].issues
    assert (by_row[5].person.first_display, by_row[5].person.last_display) == ("Yanira", "Delgado")
    assert by_row[6].person.last_display == "Ruiz Lindo"
    assert by_row[8].person.last_display == "De Jesús"
    assert by_row[9].status == "invalid" and "COMPANY_MISSING" in by_row[9].issues
    assert by_row[3].location.city == "San Pedro Garza García"


def test_specific_header_wins_over_generic_word() -> None:
    proposed = propose_mapping(["Nombre del contacto", "Apellidos", "Nombre de la empresa"])
    assert proposed.mapping["first_name"] == "Nombre del contacto"
    assert proposed.mapping["company"] == "Nombre de la empresa"
    assert proposed.mapping["location"] is None


def test_validate_mapping_reports_missing_unknown_and_repeated() -> None:
    errors = validate_mapping(
        {"first_name": "A", "last_name": "A", "company": "Z", "title": None, "location": None},
        ["A", "B"],
    )
    assert "La columna 'Z' no existe en el archivo." in errors
    assert "La columna 'A' está asignada a más de un campo." in errors
    assert validate_mapping({"first_name": "A"}, ["A"]) == [
        "Falta asignar la columna de Apellido.",
        "Falta asignar la columna de Empresa.",
    ]


# --- filas ---------------------------------------------------------------------------------


def _fixture_rows() -> tuple[list[str], list[tuple[int, dict[str, str]]]]:
    table = read_table(FIXTURE.read_bytes(), "contactos.csv", LIMITS)
    rows = [(i, dict(zip(table.headers, r, strict=True))) for i, r in enumerate(table.rows, 2)]
    return table.headers, rows


def test_process_fixture_statuses_and_issues() -> None:
    headers, rows = _fixture_rows()
    processed, summary = process_rows(rows, propose_mapping(headers).mapping)
    by_row = {r.row_number: r for r in processed}

    assert (summary.total, summary.valid, summary.invalid, summary.duplicate) == (9, 6, 2, 1)
    assert by_row[2].status == "valid" and by_row[2].issues == ()
    assert "NAME_CLEANED" in by_row[3].issues  # "Ing." retirado
    assert by_row[5].issues == ("TITLE_MISSING",)
    assert by_row[6].status == "invalid" and by_row[6].issues == ("COMPANY_INVALID",)
    assert by_row[7].status == "invalid" and "FIRST_NAME_MISSING" in by_row[7].issues
    assert "LOCATION_UNRECOGNIZED" in by_row[8].issues
    # Misma persona, empresa y país que la fila 2 con otra grafía: duplicada.
    assert by_row[9].status == "duplicate" and by_row[9].duplicate_of == 2
    assert "LOCATION_AMBIGUOUS" in by_row[10].issues
    # Bancolombia en Colombia (filas 2 y 9) y Bancolombia sin país (fila 8) son empresas distintas.
    assert summary.companies == 6
    assert summary.countries["CO"] == 2


def test_normalized_payload() -> None:
    headers, rows = _fixture_rows()
    processed, _ = process_rows(rows[:1], propose_mapping(headers).mapping)
    assert processed[0].normalized_payload() == {
        "first_name": "María José",
        "last_name": "De la Cruz Pérez",
        "title": "Gerente de Mercadeo",
        "company": "Bancolombia",
        "company_normalized": "bancolombia",
        "country_code": "CO",
        "country": "Colombia",
        "city": "Bogotá",
        "location_confidence": "high",
    }


def test_location_issues_only_when_location_is_mapped() -> None:
    raw = {"N": "Ana", "A": "López", "E": "Rappi"}
    processed, _ = process_rows([(2, raw)], {"first_name": "N", "last_name": "A", "company": "E"})
    assert processed[0].issues == ()


# --- exportación ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("value", "expected"),
    [
        ("=1+1", "'=1+1"),
        ("+57300", "'+57300"),
        ("-x", "'-x"),
        ("@cmd", "'@cmd"),
        ("Ana", "Ana"),
        (None, ""),
    ],
)
def test_sanitize_cell(value: object, expected: str) -> None:
    assert sanitize_cell(value) == expected


def _export_rows() -> tuple[list[str], list[ExportRow]]:
    headers, rows = _fixture_rows()
    processed, _ = process_rows(rows, propose_mapping(headers).mapping)
    return headers, [
        ExportRow(
            row_number=r.row_number,
            raw=r.raw,
            status=r.status,
            issues=r.issues,
            duplicate_of=r.duplicate_of,
            normalized=r.normalized_payload(),
        )
        for r in processed
    ]


def test_csv_export_preserves_original_columns_and_neutralizes_formulas() -> None:
    headers, rows = _export_rows()
    content = to_csv(headers, rows, MESSAGES)
    assert content.startswith(b"\xef\xbb\xbf")
    records = list(csv.reader(io.StringIO(content.decode("utf-8-sig"))))
    assert records[0][:5] == headers
    assert records[0][5] == "ORKESTA estado"
    assert records[1][:5] == [
        "María José",
        "De la Cruz Pérez",
        "Gerente de Mercadeo",
        "Bancolombia S.A.",
        "Bogotá, Distrito Capital, Colombia",
    ]
    assert records[1][5] == "válida"
    assert records[8][5] == "duplicada" and records[8][7] == "2"
    # La fila con =HYPERLINK(...) se exporta como texto.
    assert records[9][2].startswith("'=HYPERLINK")


def test_xlsx_export_writes_formulas_as_plain_text() -> None:
    headers, rows = _export_rows()
    workbook = load_workbook(io.BytesIO(to_xlsx(headers, rows, MESSAGES)))
    sheet = workbook.active
    assert sheet is not None
    assert [c.value for c in sheet[1]][:6] == [*headers, "ORKESTA estado"]
    cell = sheet.cell(row=10, column=3)
    assert cell.value == '=HYPERLINK("http://ejemplo.invalid")'
    assert cell.data_type == "s"
