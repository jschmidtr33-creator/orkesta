"""Listas de prospección: correos principal + alternativos, miembros, interfaz y exportación."""

from __future__ import annotations

import io
import re
import uuid
from collections.abc import Iterator

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import Engine, select
from sqlalchemy.orm import Session, sessionmaker

from orkesta.api.main import create_app
from orkesta.auth import create_user
from orkesta.catalog.lists import (
    ListError,
    add_to_list,
    compute_emails,
    create_list,
    refresh_emails,
    remove_members,
)
from orkesta.catalog.prospects import records_from_file
from orkesta.catalog.service import refresh_domain_patterns, upsert_prospects
from orkesta.db.models import Organization, Prospect, ProspectListMember

pytestmark = pytest.mark.db

BASE = (
    "email,first_name,last_name,title,company,country\n"
    "ana.lopez@lacteos.com.mx,Ana,López,Directora de Compras,Lácteos Inventados,Mexico\n"
    "jorge.diaz@lacteos.com.mx,Jorge,Díaz,CIO,Lácteos Inventados,Mexico\n"
    ",Carla,Méndez,CFO,Lácteos Inventados,Mexico\n"
    "luis@gmail.com,Luis,Ruiz,CEO,Sin Dominio SA,Argentina\n"
    "mperez@bancoficticio.com.ar,María,Pérez,Jefa de TI,Banco Ficticio,Argentina\n"
).encode()


@pytest.fixture()
def org_id(session_factory: sessionmaker[Session]) -> uuid.UUID:
    with session_factory.begin() as session:
        org = Organization(name="Org", slug=f"org-{uuid.uuid4().hex[:8]}")
        session.add(org)
        session.flush()
        upsert_prospects(session, org.id, records_from_file(BASE, "b.csv"), source="base")
        refresh_domain_patterns(session, org.id, ["lacteos.com.mx"])
        return org.id


def _people(session: Session, org_id: uuid.UUID) -> dict[str, Prospect]:
    rows = session.scalars(select(Prospect).where(Prospect.organization_id == org_id))
    return {p.first_name: p for p in rows}


def test_emails_known_real_pattern_and_missing_domain(
    session_factory: sessionmaker[Session], org_id: uuid.UUID
) -> None:
    with session_factory() as session:
        people = _people(session, org_id)
        emails = compute_emails(session, org_id, list(people.values()))
    ana = emails[people["Ana"].id]
    assert (ana.email, ana.confidence, ana.source) == ("ana.lopez@lacteos.com.mx", 100, "known")
    assert ana.alt_1 and ana.alt_1 != ana.email and ana.alt_1.endswith("@lacteos.com.mx")
    # Sin correo propio: dominio de sus colegas y patrón real de la empresa.
    carla = emails[people["Carla"].id]
    assert carla.email == "carla.mendez@lacteos.com.mx" and carla.source == "real_pattern"
    assert carla.alt_1 and carla.alt_2
    # Correo conocido en dominio corporativo sin patrón aprendido aún: se conserva el conocido.
    maria = emails[people["María"].id]
    assert maria.email == "mperez@bancoficticio.com.ar" and maria.source == "known"
    # Solo correo gratuito y sin colegas con dominio: el correo gratuito se conserva.
    luis = emails[people["Luis"].id]
    assert luis.email == "luis@gmail.com" and luis.alt_1 is None


def test_add_remove_exclude_and_refresh(
    session_factory: sessionmaker[Session], org_id: uuid.UUID
) -> None:
    with session_factory.begin() as session:
        people = _people(session, org_id)
        first = create_list(session, organization_id=org_id, user_id=None, name="CIOs México")
        second = create_list(session, organization_id=org_id, user_id=None, name="Todos")
        with pytest.raises(ListError):
            create_list(session, organization_id=org_id, user_id=None, name="cios  méxico")
        ids = [people["Ana"].id, people["Carla"].id]
        assert (
            add_to_list(session, organization_id=org_id, list_id=first.id, prospect_ids=ids)[
                "agregadas"
            ]
            == 2
        )
        again = add_to_list(session, organization_id=org_id, list_id=first.id, prospect_ids=ids)
        assert (again["agregadas"], again["ya_estaban"]) == (0, 2)
        everyone = [p.id for p in people.values()]
        result = add_to_list(
            session,
            organization_id=org_id,
            list_id=second.id,
            prospect_ids=everyone,
            exclude_in_other_lists=True,
        )
        assert (result["agregadas"], result["en_otras_listas"]) == (3, 2)
        member = session.scalar(
            select(ProspectListMember).where(
                ProspectListMember.list_id == first.id,
                ProspectListMember.prospect_id == people["Carla"].id,
            )
        )
        assert member is not None and member.email == "carla.mendez@lacteos.com.mx"
        assert remove_members(session, org_id, first.id, [member.id]) == 1
        second_id = second.id
    # Nueva base con colegas de Banco Ficticio que enseñan su patrón: recalcular actualiza.
    extra = (
        b"email,first_name,last_name,company,country\n"
        b"jgomez@bancoficticio.com.ar,Juan,Gomez,Banco Ficticio,Argentina\n"
        b",Sofia,Rios,Banco Ficticio,Argentina\n"
    )
    with session_factory.begin() as session:
        upsert_prospects(session, org_id, records_from_file(extra, "e.csv"), source="base")
        refresh_domain_patterns(session, org_id, ["bancoficticio.com.ar"])
        sofia = session.scalar(select(Prospect).where(Prospect.first_name == "Sofia"))
        assert sofia is not None
        add_to_list(session, organization_id=org_id, list_id=second_id, prospect_ids=[sofia.id])
        assert refresh_emails(session, org_id, second_id) == 4
        member = session.scalar(
            select(ProspectListMember).where(ProspectListMember.prospect_id == sofia.id)
        )
        assert member is not None
        assert (member.email, member.email_source) == ("srios@bancoficticio.com.ar", "real_pattern")


@pytest.fixture()
def client(
    engine: Engine, session_factory: sessionmaker[Session], org_id: uuid.UUID
) -> Iterator[TestClient]:
    with session_factory.begin() as session:
        create_user(
            session, email="jenn@ejemplo.com", password="clave-de-prueba-1", organization_id=org_id
        )
    with TestClient(create_app(engine_provider=lambda: engine), follow_redirects=False) as c:
        token = re.search(r'name="csrf_token" value="([^"]+)"', c.get("/login").text)
        assert token
        c.post(
            "/login",
            data={
                "email": "jenn@ejemplo.com",
                "password": "clave-de-prueba-1",
                "csrf_token": token.group(1),
            },
        )
        yield c


def _token(client: TestClient, path: str) -> str:
    match = re.search(r'name="csrf_token" value="([^"]+)"', client.get(path).text)
    assert match
    return match.group(1)


def test_lists_from_search_to_export(client: TestClient) -> None:
    token = _token(client, "/prospectos")
    # Todos los decisores de México a una lista nueva, desde el buscador.
    response = client.post(
        "/prospectos/a-lista",
        data={
            "csrf_token": token,
            "nueva_lista": "Decisores México",
            "alcance": "todos",
            "pais": "MX",
            "nivel": ["c_level", "duenio_socio", "vp", "director", "jefe"],
        },
    )
    assert response.status_code == 303
    lists_page = client.get("/listas").text
    assert "Decisores México" in lists_page
    list_url = re.search(r'href="(/listas/[0-9a-f-]{36})"', lists_page)
    assert list_url
    detail = client.get(list_url.group(1)).text
    assert "carla.mendez@lacteos.com.mx" in detail and "patrón real de la empresa" in detail
    assert "Luis Ruiz" not in detail  # Argentina
    xlsx = client.get(f"{list_url.group(1)}/exportar", params={"formato": "xlsx"})
    sheet = load_workbook(io.BytesIO(xlsx.content)).active
    assert sheet is not None
    header = [c.value for c in next(sheet.iter_rows(max_row=1))]
    assert header[:6] == [
        "Nombre completo", "Cargo", "Empresa", "Correo", "Confianza correo", "Origen correo",
    ]  # fmt: skip
    assert sheet.max_row == 4  # Ana, Jorge y Carla
    csv_text = client.get(f"{list_url.group(1)}/exportar", params={"formato": "csv"}).text
    assert "Correo alternativo 1" in csv_text
    # Eliminar la lista no borra a las personas.
    token = _token(client, list_url.group(1))
    assert (
        client.post(f"{list_url.group(1)}/eliminar", data={"csrf_token": token}).status_code == 303
    )
    assert f'href="{list_url.group(1)}"' not in client.get("/listas").text
    assert "Carla Méndez" in client.get("/prospectos", params={"cargo": "cfo"}).text


def test_other_organizations_cannot_see_a_list(
    client: TestClient, session_factory: sessionmaker[Session]
) -> None:
    with session_factory.begin() as session:
        other = Organization(name="Otra", slug=f"otra-{uuid.uuid4().hex[:8]}")
        session.add(other)
        session.flush()
        foreign = create_list(session, organization_id=other.id, user_id=None, name="Ajena")
        foreign_id = foreign.id
    assert client.get(f"/listas/{foreign_id}").status_code == 404
    assert client.get(f"/listas/{foreign_id}/exportar").status_code == 404


def test_home_dashboard(client: TestClient) -> None:
    page = client.get("/")
    assert page.status_code == 200
    assert "Decisores en tu base" in page.text and "Buscar decisores" in page.text
    # Ana, Jorge, Carla, Luis y María son decisores (directora, CIO, CFO, CEO, jefa).
    assert re.search(
        r'stat-value">5</span><span class="stat-label">Decisores en tu base', page.text
    )
